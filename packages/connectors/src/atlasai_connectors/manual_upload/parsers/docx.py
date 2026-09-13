"""DOCX parsing via python-docx. Heading paragraphs build a section_path
(e.g. "Scope > Phase 1 > Deliverables") that following body paragraphs are
tagged with, so a citation can point at "the paragraph under this heading"
rather than just an opaque chunk index.
"""

from __future__ import annotations

import io

import docx

from atlasai_connectors.manual_upload.chunking import chunk_text
from atlasai_connectors.manual_upload.parsers.base import ParsedChunk, ParsedDocument

_HEADING_STYLE_PREFIX = "Heading"


def parse_docx(data: bytes) -> ParsedDocument:
    document = docx.Document(io.BytesIO(data))

    heading_stack: list[tuple[int, str]] = []  # (level, text)
    section_buffer: list[str] = []
    chunks: list[ParsedChunk] = []
    full_text_parts: list[str] = []

    def flush_section() -> None:
        if not section_buffer:
            return
        text = "\n\n".join(section_buffer)
        full_text_parts.append(text)
        section_path = " > ".join(h[1] for h in heading_stack) or None
        for piece in chunk_text(text):
            chunks.append(ParsedChunk(content=piece, section_path=section_path))
        section_buffer.clear()

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.strip()
        if not text:
            continue

        if style_name.startswith(_HEADING_STYLE_PREFIX):
            flush_section()
            try:
                level = int(style_name.rsplit(" ", 1)[-1])
            except ValueError:
                level = 1
            heading_stack[:] = [h for h in heading_stack if h[0] < level]
            heading_stack.append((level, text))
            full_text_parts.append(text)
        else:
            section_buffer.append(text)

    flush_section()

    return ParsedDocument(full_text="\n\n".join(full_text_parts), chunks=chunks)
